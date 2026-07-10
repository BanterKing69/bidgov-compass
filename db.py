"""SQLite persistence + flat-file export for the normalised tender store."""

from __future__ import annotations

import csv
import sqlite3
from pathlib import Path
from typing import Iterable, Optional

from schema import Tender, COLUMNS, CREATE_TABLE_SQL, POST_V1_COLUMNS, POST_V1_INDEXES

DATA_DIR = Path(__file__).resolve().parent / "data"
DB_PATH = DATA_DIR / "tenders.db"
CSV_PATH = DATA_DIR / "tenders.csv"


def _ensure_columns(conn: sqlite3.Connection) -> None:
    """Idempotent live migration: adds any POST_V1_COLUMNS the current table
    doesn't have (SQLite lacks IF NOT EXISTS on ALTER TABLE ADD COLUMN, so we
    check first via PRAGMA), then creates their indexes once the columns
    exist."""
    existing = {r[1] for r in conn.execute("PRAGMA table_info(tenders)")}
    for name, coltype in POST_V1_COLUMNS:
        if name not in existing:
            conn.execute(f"ALTER TABLE tenders ADD COLUMN {name} {coltype}")
    for stmt in POST_V1_INDEXES:
        conn.execute(stmt)
    conn.commit()


def connect(db_path: Path = DB_PATH) -> sqlite3.Connection:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.executescript(CREATE_TABLE_SQL)
    _ensure_columns(conn)
    return conn


def upsert(conn: sqlite3.Connection, tenders: Iterable[Tender]) -> int:
    """Insert-or-replace by uid. Returns number of rows written."""
    placeholders = ",".join(["?"] * len(COLUMNS))
    sql = f"INSERT OR REPLACE INTO tenders ({','.join(COLUMNS)}) VALUES ({placeholders})"
    rows = [[getattr(t, c) for c in COLUMNS] for t in tenders]
    conn.executemany(sql, rows)
    conn.commit()
    return len(rows)


def export_csv(conn: sqlite3.Connection, csv_path: Path = CSV_PATH,
               open_only: bool = False) -> int:
    """Dump the table to CSV (raw_json omitted from CSV for readability)."""
    cols = [c for c in COLUMNS if c != "raw_json"]
    where = "WHERE is_open = 1" if open_only else ""
    cur = conn.execute(
        f"SELECT {','.join(cols)} FROM tenders {where} "
        f"ORDER BY (deadline IS NULL), deadline ASC"
    )
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    n = 0
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for row in cur:
            w.writerow(row)
            n += 1
    return n


def export_xlsx(conn: sqlite3.Connection, xlsx_path: Path | None = None,
                open_only: bool = True) -> Optional[Path]:
    """Optional business-friendly Excel export (skips silently if openpyxl absent)."""
    try:
        from openpyxl import Workbook
    except ImportError:
        return None
    xlsx_path = xlsx_path or (DATA_DIR / "tenders.xlsx")
    cols = [c for c in COLUMNS if c != "raw_json"]
    where = "WHERE is_open = 1" if open_only else ""
    cur = conn.execute(
        f"SELECT {','.join(cols)} FROM tenders {where} "
        f"ORDER BY (deadline IS NULL), deadline ASC"
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Open Tenders"
    ws.append(cols)
    for row in cur:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    return xlsx_path


def recategorize(conn: sqlite3.Connection) -> int:
    """Re-apply the current CPV->category map to every existing row.

    Useful after expanding CPV_CATEGORY_MAP: rows collected before the map grew
    get their category refreshed without re-scraping.
    """
    from schema import map_cpv_to_category
    cur = conn.execute("SELECT uid, cpv_code, category FROM tenders")
    updates = []
    for uid, cpv, old_cat in cur:
        new_cat = map_cpv_to_category(cpv)
        if new_cat != old_cat:
            updates.append((new_cat, uid))
    if updates:
        conn.executemany("UPDATE tenders SET category = ? WHERE uid = ?", updates)
        conn.commit()
    return len(updates)


def stats(conn: sqlite3.Connection) -> dict:
    def one(sql, *a):
        return conn.execute(sql, a).fetchone()[0]

    return {
        "total": one("SELECT COUNT(*) FROM tenders"),
        "open": one("SELECT COUNT(*) FROM tenders WHERE is_open = 1"),
        "categorised_open": one(
            "SELECT COUNT(*) FROM tenders WHERE is_open = 1 AND category IS NOT NULL"
        ),
        "by_source": conn.execute(
            "SELECT source, COUNT(*) FROM tenders GROUP BY source ORDER BY 2 DESC"
        ).fetchall(),
        "by_category": conn.execute(
            "SELECT COALESCE(category,'(unmapped)'), COUNT(*) FROM tenders "
            "WHERE is_open = 1 GROUP BY 1 ORDER BY 2 DESC"
        ).fetchall(),
    }
